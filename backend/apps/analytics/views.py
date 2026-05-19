"""
Module 6: Analytics & Reporting Views.
Compiles stats, generates audit reports, and exports PDF/Excel/CSV dashboards.
Satisfies reportlab and openpyxl checklist items.
"""
import csv
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Count, Sum, Avg, Q
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated

from apps.users.models import User
from apps.virtual_accounts.models import VirtualAccount, PaymentStatus
from apps.evaluation.models import SelfEvaluation, PromotionCycle
from apps.superior_eval.models import SuperiorEvaluation
from apps.promotion.models import PromotionDecision
from apps.evaluation.views import get_active_cycle
from core.permissions import IsAdmin, IsSuperior
from core.audit import AuditLog

# ReportLab imports for PDF generation
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

# OpenPyXL imports for Excel generation
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


class AnalyticsDashboardView(APIView):
    """
    Returns high-level statistics for Admin/Superiors:
    - Payments collected vs target
    - Evaluations status
    - Eligible vs Ineligible counts
    - Recommendations breakdown
    """
    permission_classes = [IsSuperior]

    def get(self, request):
        cycle = get_active_cycle()
        cycle_year = cycle.year if cycle else str(timezone.now().year)

        # 1. Payment stats
        total_registered_staff = User.objects.filter(role="staff", is_active=True).count()
        payments = PaymentStatus.objects.filter(cycle=cycle_year)
        total_paid_count = payments.filter(paid=True).count()
        total_revenue = VirtualAccount.objects.filter(
            cycle=cycle_year, status="confirmed"
        ).aggregate(sum=Sum("amount_due"))["sum"] or 0.0

        # 2. Evaluation status
        total_self_evals = SelfEvaluation.objects.filter(cycle=cycle_year, is_submitted=True).count()
        total_superior_evals = SuperiorEvaluation.objects.filter(cycle=cycle_year, is_submitted=True).count()

        # 3. Decision eligibility & recommendations
        decisions = PromotionDecision.objects.filter(cycle=cycle_year)
        elig_stats = decisions.aggregate(
            eligible=Count("id", filter=Q(is_eligible_by_maturity=True)),
            ineligible=Count("id", filter=Q(is_eligible_by_maturity=False)),
            rec_promote=Count("id", filter=Q(recommendation="promote")),
            rec_retain=Count("id", filter=Q(recommendation="retain")),
            rec_defer=Count("id", filter=Q(recommendation="defer")),
            rec_pending=Count("id", filter=Q(recommendation="pending")),
        )

        # 4. MDA Breakdown
        mda_stats = decisions.values("staff__mda").annotate(
            count=Count("id"),
            avg_score=Avg("average_score"),
            promoted=Count("id", filter=Q(recommendation="promote"))
        ).order_by("-count")

        data = {
            "cycle": cycle_year,
            "staff_count": total_registered_staff,
            "payments": {
                "paid_count": total_paid_count,
                "unpaid_count": max(0, total_registered_staff - total_paid_count),
                "total_revenue": float(total_revenue),
            },
            "evaluations": {
                "self_submitted": total_self_evals,
                "superior_submitted": total_superior_evals,
                "pending_forms": max(0, total_registered_staff - total_self_evals),
            },
            "eligibility": {
                "eligible": elig_stats["eligible"] or 0,
                "ineligible": elig_stats["ineligible"] or 0,
            },
            "recommendations": {
                "promote": elig_stats["rec_promote"] or 0,
                "retain": elig_stats["rec_retain"] or 0,
                "defer": elig_stats["rec_defer"] or 0,
                "pending": elig_stats["rec_pending"] or 0,
            },
            "mda_breakdown": list(mda_stats)
        }

        return Response(data)


class AuditTrailListView(APIView):
    """View system audit logs for Admins."""
    permission_classes = [IsAdmin]

    def get(self, request):
        logs = AuditLog.objects.select_related("user").all()[:150]
        data = []
        for log in logs:
            data.append({
                "id": log.id,
                "user_email": log.user.email if log.user else "System",
                "action": log.get_action_display(),
                "target_model": log.target_model,
                "target_id": log.target_id,
                "ip_address": log.ip_address,
                "timestamp": log.timestamp.isoformat(),
                "extra": log.extra,
            })
        return Response(data)


# ─── EXPORTING REPORTS ───────────────────────────────────────────────────────

class ExportCandidatesCSVView(APIView):
    """Generates standard CSV list of candidates for direct download."""
    permission_classes = [IsSuperior]

    def get(self, request):
        cycle = get_active_cycle()
        cycle_year = cycle.year if cycle else str(timezone.now().year)

        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = f'attachment; filename="promotion_candidates_{cycle_year}.csv"'

        writer = csv.writer(response)
        writer.writerow([
            "Staff Email", "Staff Name", "MDA", "Grade Level",
            "Maturity Years", "Eligible By Maturity",
            "Self score", "Superior score", "Average score", "Decision"
        ])

        decisions = PromotionDecision.objects.select_related("staff").filter(cycle=cycle_year)
        for d in decisions:
            writer.writerow([
                d.staff.email,
                d.staff.full_name,
                d.staff.mda,
                d.grade_level_at_eval,
                d.years_in_post,
                "YES" if d.is_eligible_by_maturity else "NO",
                d.self_score,
                d.superior_score,
                d.average_score,
                d.get_recommendation_display()
            ])

        return response


class ExportCandidatesExcelView(APIView):
    """Generates beautifully styled Excel spreadsheet of candidates (openpyxl)."""
    permission_classes = [IsSuperior]

    def get(self, request):
        cycle = get_active_cycle()
        cycle_year = cycle.year if cycle else str(timezone.now().year)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Promotion Candidates"

        # Sheet styling tokens
        font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        font_body = Font(name="Segoe UI", size=10)
        fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        thin_side = Side(border_style="thin", color="D9D9D9")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        headers = [
            "Staff Email", "Staff Name", "MDA", "Grade Level",
            "Maturity Years", "Eligible By Maturity",
            "Self Score", "Superior Score", "Average Score", "Recommendation Status"
        ]

        # Write header row
        for col_idx, text in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=text)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border
        
        ws.row_dimensions[1].height = 25

        # Write data rows
        decisions = PromotionDecision.objects.select_related("staff").filter(cycle=cycle_year)
        row_idx = 2
        for d in decisions:
            row_data = [
                d.staff.email,
                d.staff.full_name,
                d.staff.mda,
                d.grade_level_at_eval,
                d.years_in_post,
                "Eligible" if d.is_eligible_by_maturity else "Ineligible",
                d.self_score,
                d.superior_score,
                d.average_score,
                d.get_recommendation_display()
            ]

            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_body
                cell.border = thin_border
                
                # Zebra striping
                if row_idx % 2 == 0:
                    cell.fill = fill_zebra
                
                # Alignments
                if col_idx in (4, 5, 6, 7, 8, 9, 10):
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="promotion_candidates_{cycle_year}.xlsx"'
        wb.save(response)
        return response


class ExportCandidateScorecardPDFView(APIView):
    """Generates a professional performance evaluation scorecard PDF for a candidate."""
    permission_classes = [IsSuperior]

    def get(self, request, decision_id):
        try:
            d = PromotionDecision.objects.select_related("staff", "decided_by").get(pk=decision_id)
        except PromotionDecision.DoesNotExist:
            return HttpResponse("Candidate record not found.", status=404)

        response = HttpResponse(content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="Scorecard_{d.staff.full_name.replace(" ", "_")}.pdf"'

        doc = SimpleDocTemplate(
            response,
            pagesize=letter,
            rightMargin=0.5*inch, leftMargin=0.5*inch,
            topMargin=0.5*inch, bottomMargin=0.5*inch
        )

        styles = getSampleStyleSheet()
        
        # Define clean premium reportlab styling
        title_style = ParagraphStyle(
            "PDFTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=22,
            textColor=colors.HexColor("#1F497D"),
            spaceAfter=15,
            alignment=1
        )
        
        section_heading = ParagraphStyle(
            "SectionHeading",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=14,
            textColor=colors.HexColor("#16365C"),
            spaceBefore=10,
            spaceAfter=10,
            borderPadding=2
        )

        body_style = ParagraphStyle(
            "PDFBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#333333")
        )

        bold_body = ParagraphStyle(
            "PDFBodyBold",
            parent=body_style,
            fontName="Helvetica-Bold"
        )

        story = []

        # Title
        story.append(Paragraph("APPERS — STAFF PROMOTION SCORECARD", title_style))
        story.append(Paragraph(f"Cycle: {d.cycle} | Generated on: {timezone.now():%d-%b-%Y %H:%M}", body_style))
        story.append(Spacer(1, 15))

        # Personal information table
        story.append(Paragraph("Candidate Information", section_heading))
        personal_data = [
            [Paragraph("Full Name", bold_body), Paragraph(d.staff.full_name, body_style),
             Paragraph("Email Address", bold_body), Paragraph(d.staff.email, body_style)],
            [Paragraph("MDA / Agency", bold_body), Paragraph(d.staff.mda, body_style),
             Paragraph("Department", bold_body), Paragraph(d.staff.department, body_style)],
            [Paragraph("Current Grade Level", bold_body), Paragraph(d.grade_level_at_eval, body_style),
             Paragraph("Years in Post", bold_body), Paragraph(f"{d.years_in_post:.2f} years", body_style)],
            [Paragraph("Maturity Eligibility", bold_body),
             Paragraph("ELIGIBLE" if d.is_eligible_by_maturity else "INELIGIBLE", bold_body),
             Paragraph("Civil Service Route", bold_body), Paragraph("CSC" if d.is_csc_route else "MDA Direct", body_style)],
        ]
        
        t_personal = Table(personal_data, colWidths=[1.8*inch, 1.95*inch, 1.8*inch, 1.95*inch])
        t_personal.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#F2F5F9")),
            ("ALIGN", (0,0), (-1,-1), "LEFT"),
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
            ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor("#D9D9D9")),
            ("TOPPADDING", (0,0), (-1,-1), 6),
            ("BOTTOMPADDING", (0,0), (-1,-1), 6),
        ]))
        story.append(t_personal)
        story.append(Spacer(1, 15))

        # Performance Scores
        story.append(Paragraph("Performance & Assessment Metrics", section_heading))
        scores_data = [
            [Paragraph("Assessment Phase", bold_body), Paragraph("Weight/Average Score", bold_body), Paragraph("Status", bold_body)],
            [Paragraph("Self-Evaluation Score", body_style), Paragraph(f"{d.self_score:.2f} / 5.0", body_style), Paragraph("Submitted", body_style)],
            [Paragraph("Superior Evaluation Score", body_style), Paragraph(f"{d.superior_score:.2f} / 5.0", body_style), Paragraph("Submitted", body_style)],
            [Paragraph("Combined Average Score", bold_body), Paragraph(f"{d.average_score:.2f} / 5.0", bold_body), Paragraph("Calculated", bold_body)],
        ]
        t_scores = Table(scores_data, colWidths=[3*inch, 2.5*inch, 2*inch])
        t_scores.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1F497D")),
            ("TEXTCOLOR", (0,0), (-1,0), colors.white),
            ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor("#D9D9D9")),
            ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F9F9F9")]),
            ("TOPPADDING", (0,0), (-1,-1), 6),
            ("BOTTOMPADDING", (0,0), (-1,-1), 6),
        ]))
        
        # Override headers color to white paragraph style
        white_bold = ParagraphStyle("WhiteBold", parent=bold_body, textColor=colors.white)
        scores_data[0] = [Paragraph("Assessment Phase", white_bold), Paragraph("Weight/Average Score", white_bold), Paragraph("Status", white_bold)]
        
        story.append(t_scores)
        story.append(Spacer(1, 15))

        # Decision & Approvals
        story.append(Paragraph("Official Promotion Panel Decision", section_heading))
        rec_str = d.get_recommendation_display().upper() if d.recommendation != "pending" else "PENDING ACTION"
        decision_data = [
            [Paragraph("Final Recommendation", bold_body), Paragraph(rec_str, bold_body)],
            [Paragraph("Panel Justification / Remarks", bold_body), Paragraph(d.remarks or "No remarks provided.", body_style)],
            [Paragraph("Authorized Decision Maker", bold_body), Paragraph(d.decided_by_name or "Pending Panel decision.", body_style)],
            [Paragraph("External Approver Status (HoS / CSC)", bold_body), Paragraph(d.get_external_status_display().upper() if d.is_csc_route else "Not Applicable", body_style)],
        ]
        t_decision = Table(decision_data, colWidths=[2.5*inch, 5*inch])
        t_decision.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#F9EBEA") if d.recommendation == "defer" else colors.HexColor("#EAFAF1")),
            ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor("#D9D9D9")),
            ("TOPPADDING", (0,0), (-1,-1), 8),
            ("BOTTOMPADDING", (0,0), (-1,-1), 8),
            ("VALIGN", (0,0), (-1,-1), "TOP"),
        ]))
        story.append(t_decision)
        story.append(Spacer(1, 30))

        # Signatures
        story.append(Paragraph("___________________________<br/><b>Director Human Resources</b>", body_style))

        doc.build(story)
        return response
